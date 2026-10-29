"""Read-only, loopback-only mock of Protheus read APIs backed by Excel reports.

This tool never connects to the application database and has no write routes.
The response contract is a test fixture, not a claim about the real Protheus API.
"""

from __future__ import annotations

import argparse
import csv
import json
import threading
import time
import uuid
from datetime import date, datetime
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from itertools import islice
from pathlib import Path
from typing import Any
from urllib.parse import parse_qs, urlparse

try:
    from openpyxl import load_workbook
except ImportError as exc:  # pragma: no cover - environment setup message
    raise SystemExit("Dependência ausente: instale openpyxl no Python usado para executar o mock.") from exc


REPORTS = {
    "products": ("produtos.xlsx", "Listagem do Browse", 2),
    "product-groups": ("Grupos dos Produtos.xlsx", "Listagem do Browse", 2),
    "resources": ("SH1.xlsx", "scbtjs40", 4),
    "structures": ("Estrutura.xlsx", "Listagem do Browse", 2),
    "routings": ("Processo Produtivo + Tempo.xlsx", "Processo Produtivo = Tempo", 2),
    "routings-basic": ("Processo Produtivo.xlsx", "Listagem do Browse", 2),
    "production-orders": ("Ordens de Produção 17-07.xlsx", "Listagem do Browse", 2),
    "sales-order-headers": ("SC5.xlsx", "scbtjok0", 4),
    "sales-order-lines": ("SC6.xlsx", "scbtjoo0", 4),
    "appointments": ("apontamento de produção.xlsx", "Listagem do Browse", 2),
    # This export stores one CSV record in each Excel cell in column A.
    "commitments": ("Empenhos.xlsx", "scbtdt90", 4),
}


def json_value(value: Any) -> Any:
    if isinstance(value, (datetime, date)):
        return value.isoformat()
    if value is None or isinstance(value, (str, int, float, bool)):
        return value
    return str(value)


def map_row(entity: str, row: tuple[Any, ...]) -> dict[str, Any]:
    v = lambda index: json_value(row[index]) if index < len(row) else None
    if entity == "products":
        return {"productCode": v(0), "description": v(1), "type": v(2), "unit": v(3), "groupCode": v(5)}
    if entity == "product-groups":
        return {"groupCode": v(0), "description": v(1), "brand": v(2)}
    if entity == "resources":
        return {
            "branch": v(0), "resourceCode": v(1), "resourceName": v(2),
            "productionLineRaw": v(3), "laborEfficiencyRaw": v(4),
            "costCenter": v(5), "lastMaintenanceRaw": v(6),
            "maintenanceIntervalRaw": v(7), "utilizationIndexRaw": v(8),
            "workCenterRaw": v(9), "calendarRaw": v(10),
            "unlimitedResourceFlagRaw": v(11), "costCenterGroupRaw": v(12),
        }
    if entity == "structures":
        return {
            "parentProductCode": v(0), "componentCode": v(1), "componentDescription": v(2),
            "sourceSequence": v(3), "revisionEnd": v(4), "quantity": v(5), "unit": v(6),
            "phantom": v(7), "lossIndex": v(8), "revisionStart": v(10), "warehouse": v(12),
        }
    if entity == "routings":
        return {
            "branch": v(0), "routingCode": v(1), "productCode": v(2), "operationSequence": v(3),
            "resource": v(4), "setupHours": v(5), "standardTimeRaw": v(6),
            "operationDescription": v(7), "lineType": v(8), "department": v(12),
            "labor": v(13), "lotSize": v(15), "operationType": v(16),
            "alternateRouting": v(23), "validFromRaw": v(26), "validUntilRaw": v(27),
        }
    if entity == "routings-basic":
        return {
            "branch": v(0), "routingCode": v(1), "productCode": v(2),
            "operationSequence": v(3), "operationDescription": v(4),
        }
    if entity == "production-orders":
        return {
            "sourceCompositeKey": v(0), "orderNumber": v(1), "orderItem": v(2),
            "orderSequence": v(3), "productCode": v(4), "productDescription": v(5),
            "warehouse": v(6), "plannedQuantity": v(7), "forecastStartRaw": v(8),
            "dueDateRaw": v(9), "issuedAtRaw": v(10), "actualEndRaw": v(11),
            "producedQuantity": v(12), "sourceOPKey": v(13), "valueClass": v(14),
            "valueClassDescription": v(15), "sourceOrderType": v(16), "productRevision": v(17),
            # These facts are not present in this report. Do not derive or default them.
            "scrapQuantity": None, "sourceBalance": None, "salesOrderNumber": None,
        }
    if entity == "sales-order-headers":
        return {
            "branch": v(0), "orderNumber": v(1), "orderType": v(2),
            "customerCode": v(3), "customerStore": v(4),
            "deliveryCustomerCode": v(5), "deliveryCustomerStore": v(6),
            "customerName": v(7), "headerDueDateRaw": v(12), "issuedAtRaw": v(28),
            # SC5 is the order header; this SC6 export lacks C6_NUM to link lines.
            "salesOrderItem": None, "productCode": None, "orderedQuantity": None,
            "deliveredQuantity": None, "unit": None, "lineDueDateRaw": None,
            "lineStatus": None,
        }
    if entity == "sales-order-lines":
        return {
            "branch": v(0), "salesOrderItem": v(1), "productCode": v(2),
            "description": v(3), "unit": v(4), "orderedQuantity": v(5),
            "unitPriceRaw": v(6), "lineTotalRaw": v(7), "releasedQuantity": v(8),
            "deliveryDateRaw": v(11), "originProductionOrder": v(12),
            "beneficiaryProductionOrder": v(13), "blockCodeRaw": v(34),
            "customerPurchaseOrderRaw": v(32), "deliveredQuantity": v(27),
            "blockStatusRaw": v(40), "purchaseRequestNumberRaw": v(72),
            "purchaseRequestItemRaw": v(73), "salesOrderNumber": None,
        }
    if entity == "appointments":
        return {
            "orderNumber": v(0), "productCode": v(1), "operationRaw": v(2),
            "startedAtDateRaw": v(3), "startedAtTimeRaw": v(4),
            "endedAtDateRaw": v(5), "endedAtTimeRaw": v(6),
            "producedQuantity": v(7), "actualTimeRaw": v(8),
            "observationsRaw": v(9), "operator": v(10),
            "qualityCertificateRaw": v(11), "lotSizeRaw": v(12),
            "warehouseRaw": v(13), "enteredByRaw": v(14),
            "operationSequence": None, "workCenter": None, "scrapQuantity": None,
        }
    if entity == "commitments":
        values = next(csv.reader([str(row[0] or "")]))
        values += [None] * max(0, 29 - len(values))
        return {
            "branch": values[0], "productCode": values[1], "description": values[2],
            "productionOrder": values[3], "warehouse": values[4], "commitmentDateRaw": values[5],
            "suspendedQuantity": values[6], "sourceSituation": values[7],
            "committedQuantity": values[8], "commitmentBalance": values[9],
            "valueClass": values[10], "structureSequence": values[11], "lot": values[12],
            "subLot": values[13], "expirationDateRaw": values[14], "originOrder": values[15],
            "alternateUnitBalance": values[16], "maintenanceOrder": values[17],
            "sourceSequence": values[18], "lotPower": values[19], "processQuantity": values[20],
            "parentProductCode": values[21], "routingCode": values[22], "operation": values[23],
            "phantom": values[25], "observation": values[26], "createdByName": values[27],
            "createdByCode": values[28],
        }
    raise KeyError(entity)


class ReportStore:
    def __init__(self, data_dir: Path):
        self.data_dir = data_dir.expanduser().resolve()
        self._workbooks = {}
        self._cursors = {}
        self._cursor_lock = threading.Lock()
        missing = [name for name, *_ in REPORTS.values() if not (self.data_dir / name).is_file()]
        if missing:
            raise FileNotFoundError("Relatórios não encontrados: " + ", ".join(missing))

    def page(self, entity: str, page: int, page_size: int) -> dict[str, Any]:
        if entity not in REPORTS:
            raise KeyError(entity)
        filename, sheet_name, first_data_row = REPORTS[entity]
        path = self.data_dir / filename
        if filename not in self._workbooks:
            self._workbooks[filename] = load_workbook(path, read_only=True, data_only=True)
        workbook = self._workbooks[filename]
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"A aba '{sheet_name}' não existe em {filename}.")
        sheet = workbook[sheet_name]
        total = max(0, sheet.max_row - first_data_row + 1)
        start = first_data_row + (page - 1) * page_size
        end = min(first_data_row + total - 1, start + page_size - 1)
        rows = sheet.iter_rows(min_row=start, max_row=end, max_col=1 if entity == "commitments" else sheet.max_column, values_only=True)
        records = [map_row(entity, row) for row in rows if any(value is not None for value in row)]

        pages = (total + page_size - 1) // page_size
        return {
            "simulation": True,
            "readOnly": True,
            "sourceFile": filename,
            "sourceSheet": sheet_name,
            "sourceRows": total,
            "page": page,
            "pageSize": page_size,
            "totalPages": pages,
            "recordsReturned": len(records),
            "records": records,
        }

    def _record_iterator(self, entity: str):
        filename, sheet_name, first_data_row = REPORTS[entity]
        workbook = self._workbooks.get(filename)
        if workbook is None:
            workbook = load_workbook(self.data_dir / filename, read_only=True, data_only=True)
            if sheet_name not in workbook.sheetnames:
                workbook.close()
                raise ValueError(f"A aba '{sheet_name}' não existe em {filename}.")
            self._workbooks[filename] = workbook
        sheet = workbook[sheet_name]
        total = max(0, sheet.max_row - first_data_row + 1)
        max_col = 1 if entity == "commitments" else sheet.max_column
        for row in sheet.iter_rows(min_row=first_data_row, max_row=first_data_row + total - 1, max_col=max_col, values_only=True):
            if any(value is not None for value in row):
                yield map_row(entity, row)

    def cursor_page(self, entity: str, page_size: int, cursor: str | None = None, probe: bool = False) -> dict[str, Any]:
        if entity not in REPORTS:
            raise KeyError(entity)
        filename, sheet_name, first_data_row = REPORTS[entity]
        workbook = self._workbooks.get(filename)
        if workbook is None:
            workbook = load_workbook(self.data_dir / filename, read_only=True, data_only=True)
            if sheet_name not in workbook.sheetnames:
                workbook.close()
                raise ValueError(f"A aba '{sheet_name}' não existe em {filename}.")
            self._workbooks[filename] = workbook
        total = max(0, workbook[sheet_name].max_row - first_data_row + 1)

        with self._cursor_lock:
            now = time.monotonic()
            expired = [token for token, state in self._cursors.items() if state["expires"] <= now]
            for token in expired:
                self._cursors.pop(token)["iterator"].close()
            if cursor:
                state = self._cursors.pop(cursor, None)
                if not state or state["entity"] != entity:
                    raise ValueError("cursor inválido, expirado ou de outra entidade")
                iterator = state["iterator"]
                page_number = state["page"] + 1
                pending = state.get("pending")
            else:
                iterator = self._record_iterator(entity)
                page_number = 1
                pending = None

            records = []
            if pending is not None:
                records.append(pending)
            records.extend(list(islice(iterator, page_size - len(records))))
            lookahead = next(iterator, None)
            next_cursor = None
            if lookahead is not None:
                if probe:
                    iterator.close()
                else:
                    next_cursor = uuid.uuid4().hex
                    self._cursors[next_cursor] = {
                        "entity": entity, "iterator": iterator, "page": page_number,
                        "pending": lookahead, "expires": now + 600,
                    }
            else:
                iterator.close()

        return {
            "simulation": True, "readOnly": True, "sourceFile": filename,
            "sourceSheet": sheet_name, "sourceRows": total, "page": page_number,
            "pageSize": page_size, "recordsReturned": len(records),
            "nextCursor": next_cursor, "cursorProbeComplete": bool(probe and cursor), "records": records,
        }

    def close(self):
        for workbook in self._workbooks.values():
            workbook.close()
        self._workbooks.clear()


def make_handler(store: ReportStore):
    class Handler(BaseHTTPRequestHandler):
        server_version = "ProtheusMock/1.0"

        def _send(self, status: int, payload: dict[str, Any]):
            body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
            self.send_response(status)
            self.send_header("Content-Type", "application/json; charset=utf-8")
            self.send_header("Content-Length", str(len(body)))
            self.send_header("Cache-Control", "no-store")
            self.end_headers()
            self.wfile.write(body)

        def do_GET(self):
            parsed = urlparse(self.path)
            if parsed.path in ("/", "/v1"):
                return self._send(200, {
                    "name": "API Mock Protheus",
                    "simulation": True,
                    "readOnly": True,
                    "connectsToProtheus": False,
                    "writesToApplication": False,
                    "pagination": {"page": "1-based", "pageSizeMax": 500},
                    "endpoints": [f"GET /v1/{entity}" for entity in REPORTS]
                    + ["GET /v1/sales-orders (501; SC6 lacks the sales-order link key C6_NUM)", "GET /v1/health", "GET /v1/flow-check (full read-only preflight)"],
                })
            if parsed.path == "/v1/health":
                return self._send(200, {"status": "online", "simulation": True, "readOnly": True, "connectsToProtheus": False, "writesToApplication": False})
            if parsed.path == "/v1/flow-check":
                try:
                    from protheus_mock_flow import run_preflight

                    class StoreFlowClient:
                        def __init__(self, report_store):
                            self.report_store = report_store
                            self.counts = {}

                        def health(self):
                            return {"simulation": True, "readOnly": True, "connectsToProtheus": False, "writesToApplication": False}

                        def records(self, entity):
                            received = 0
                            for record in self.report_store._record_iterator(entity):
                                received += 1
                                yield record
                            self.counts[entity] = received

                    client = StoreFlowClient(store)
                    return self._send(200, run_preflight(client))
                except Exception:
                    return self._send(503, {
                        "error": "FLOW_CHECK_FAILED",
                        "message": "A verificação completa dos relatórios não pôde ser concluída.",
                        "simulation": True,
                        "readOnly": True,
                        "databaseWritePerformed": False,
                    })
            if parsed.path == "/v1/sales-orders":
                return self._send(501, {"error": "SOURCE_LINK_KEY_MISSING", "simulation": True, "message": "O SC6 está disponível como linhas, mas o relatório não contém C6_NUM para vincular cada item ao cabeçalho SC5."})
            if not parsed.path.startswith("/v1/"):
                return self._send(404, {"error": "NOT_FOUND"})

            entity = parsed.path.removeprefix("/v1/")
            query = parse_qs(parsed.query)
            try:
                page = int(query.get("page", ["1"])[0])
                page_size = int(query.get("pageSize", ["100"])[0])
                if page < 1 or page_size < 1 or page_size > 500:
                    raise ValueError("page deve ser >= 1 e pageSize deve estar entre 1 e 500.")
                if "cursor" in query or query.get("stream", ["false"])[0].lower() == "true":
                    cursor = query.get("cursor", [None])[0]
                    probe = query.get("probe", ["false"])[0].lower() == "true"
                    return self._send(200, store.cursor_page(entity, page_size, cursor, probe))
                return self._send(200, store.page(entity, page, page_size))
            except KeyError:
                return self._send(404, {"error": "ENTITY_NOT_AVAILABLE", "entity": entity})
            except ValueError as error:
                return self._send(400, {"error": "INVALID_PAGINATION", "message": str(error)})
            except Exception:
                # Keep filesystem paths and workbook internals out of API responses.
                return self._send(500, {"error": "REPORT_READ_FAILED", "entity": entity})

        def do_POST(self):
            self.send_response(405)
            self.send_header("Allow", "GET")
            self.send_header("Content-Length", "0")
            self.end_headers()

        do_PUT = do_POST
        do_PATCH = do_POST
        do_DELETE = do_POST

        def log_message(self, format, *args):
            # Do not log returned business data or source rows.
            return

    return Handler


def main():
    parser = argparse.ArgumentParser(description="Mock REST somente leitura para relatórios Protheus.")
    parser.add_argument("--data-dir", required=True, type=Path, help="Pasta local que contém os relatórios Excel.")
    parser.add_argument("--port", type=int, default=4100)
    args = parser.parse_args()
    store = ReportStore(args.data_dir)
    server = ThreadingHTTPServer(("127.0.0.1", args.port), make_handler(store))
    print(f"Mock Protheus somente leitura: http://127.0.0.1:{args.port}/v1/health")
    print("Simulação local; não consulta Protheus e não grava no banco da aplicação.")
    try:
        server.serve_forever()
    except KeyboardInterrupt:
        pass
    finally:
        server.server_close()
        store.close()


if __name__ == "__main__":
    main()
